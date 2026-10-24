from io import BytesIO
from pathlib import Path
from docx import Document
from pypdf import PdfReader

MAX_UPLOAD_BYTES = 5 * 1024 * 1024

class DocumentParseError(ValueError):
    pass

def extract_document(filename: str, content: bytes) -> str:
    if len(content) > MAX_UPLOAD_BYTES:
        raise DocumentParseError("Each document must be 5 MB or smaller.")
    extension = Path(filename).suffix.lower()
    try:
        if extension in {".txt", ".md", ".csv"}:
            return content.decode("utf-8-sig", errors="replace").strip()
        if extension == ".pdf":
            reader = PdfReader(BytesIO(content))
            text = "\n".join(page.extract_text() or "" for page in reader.pages[:35])
            if not text.strip():
                raise DocumentParseError("No selectable text was found in this PDF. Paste the text manually or use a text-based PDF.")
            return text.strip()
        if extension == ".docx":
            doc = Document(BytesIO(content))
            paragraphs = [paragraph.text.strip() for paragraph in doc.paragraphs if paragraph.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        paragraphs.append(row_text)
            return "\n".join(paragraphs).strip()
    except DocumentParseError:
        raise
    except Exception as exc:
        raise DocumentParseError(f"Could not read {filename}. Try a text-based PDF, DOCX, or TXT file.") from exc
    raise DocumentParseError("Unsupported file type. Use PDF, DOCX, TXT, or MD.")
